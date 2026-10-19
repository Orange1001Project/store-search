"""서비스형 쿼리셋(service_queryset_v1) 작성용 xlsx 템플릿을 만든다.

파이프라인 번호가 없는 이유: import_queryset_xlsx.py처럼 사람이 쿼리를 쓰기 전에 한 번 돌리는 저작 도구다.
표 내용은 src/store_search_ai/data/service_queryset.py, 설정(상황 카드·별칭·목표 개수·seed)은
configs/query/service_queryset_v1.yaml, 계획은 docs/SERVICE_QUERYSET.md.

출력 파일은 작성자가 채우는 원본이 되므로 이미 있으면 덮어쓰지 않는다(--force로만).

사용법:
    python scripts/make_service_queryset_template.py
    python scripts/make_service_queryset_template.py --dry-run
"""

from __future__ import annotations

import argparse
import logging
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from store_search_ai.data.service_queryset import (
    INPUT_COLUMNS,
    absent_candidate_counts,
    chain_candidates,
    family_gap_rows,
    guide_table,
    item_gap_candidates,
    sample_markets,
    sample_stores,
    scenario_card_table,
)
from store_search_ai.pipeline.common import load_config

logger = logging.getLogger(__name__)

INPUT_ROWS = 500
COLUMN_WIDTH_MAX = 60


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default="configs/query/service_queryset_v1.yaml")
    parser.add_argument("--benchmark-config", default="configs/benchmark/storesearch_ko_v1.yaml")
    parser.add_argument("--families", default="configs/benchmark/query_families_v1.yaml")
    parser.add_argument("--queryset-xlsx", default="data/query/queryset_final.xlsx")
    parser.add_argument("--output", default=None, help="기본: 설정의 output_path")
    parser.add_argument("--force", action="store_true", help="출력 파일이 있어도 덮어쓴다(작성 내용이 사라짐)")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(message)s")

    cfg = load_config(args.config)
    bench = load_config(args.benchmark_config)
    output = Path(args.output or cfg["output_path"])
    if output.exists() and not (args.force or args.dry_run):
        raise SystemExit(f"{output}이(가) 이미 있습니다. 작성 중인 파일일 수 있어 덮어쓰지 않습니다(--force).")

    seed = cfg["random_seed"]
    corpus = pd.read_parquet(bench["corpus_path"])
    families = load_config(args.families)["families"]
    inactive = load_config(bench["query_set"]["inactive_queries_path"])["inactive_queries"]
    inactive_ids = {row["query_id"] for row in inactive}
    fam_list = pd.read_excel(args.queryset_xlsx, sheet_name="패밀리목록")

    gap_cfg = cfg["item_gap"]
    generic_items = set(gap_cfg["generic_items"])
    covered_terms = {t for f in families for t in [f["family"], *f["positive_terms"]]}
    item_gaps = item_gap_candidates(corpus, covered_terms, generic_items, gap_cfg["top_n"])
    fam_gaps = family_gap_rows(fam_list, families, inactive_ids, gap_cfg["thin_family_max_queries"], 1)
    a_candidates = pd.concat([fam_gaps, item_gaps], ignore_index=True)
    a_candidates["후보ID"] = [f"A{i + 1:03d}" for i in range(len(a_candidates))]

    chain_cfg = cfg["chain"]
    chains = chain_candidates(
        corpus, chain_cfg["aliases"], chain_cfg["generic_name_tokens"], chain_cfg["common_name_suffixes"],
        chain_cfg["min_count"],
    )
    stores = sample_stores(
        corpus, cfg["store_sample"]["n"], cfg["store_sample"]["name_only_share"], generic_items, seed
    )
    m_cfg = cfg["market_sample"]
    markets = sample_markets(corpus, m_cfg["n"], m_cfg["min_per_region"], m_cfg["min_stores"], seed)
    absent = absent_candidate_counts(corpus, cfg["absent_candidates"])
    cards = scenario_card_table(cfg["scenario_cards"])

    sheets = {
        "작성요령": guide_table(cfg["targets"], cfg["version"], seed),
        "쿼리입력": pd.DataFrame(columns=INPUT_COLUMNS),
        "A_업종후보": a_candidates,
        "B_가맹점표본": stores,
        "B_체인목록": chains,
        "C_상황카드": cards,
        "D_시장표본": markets,
        "E_결과없음후보": absent,
    }
    for name, df in sheets.items():
        logger.info("%-12s %4d행", name, len(df))
    logger.info("A 후보 구분: %s", a_candidates["구분"].value_counts().to_dict())
    logger.info("B 표본 층: %s", stores["층"].value_counts().to_dict())
    logger.info("체인 구분: %s", chains["구분"].value_counts().to_dict())
    logger.info("E 사용 가능: %d / %d", (absent["판단"] == "E로 사용 가능").sum(), len(absent))
    if args.dry_run:
        logger.info("[dry-run] 파일을 쓰지 않았습니다.")
        return

    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
        for name, df in sheets.items():
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            for idx, col in enumerate(df.columns, start=1):
                longest = max([len(str(col)), *(len(str(v)) for v in df[col].head(200))])
                ws.column_dimensions[get_column_letter(idx)].width = min(COLUMN_WIDTH_MAX, longest + 4)
        ws = writer.sheets["쿼리입력"]
        for idx in range(1, len(INPUT_COLUMNS) + 1):
            ws.column_dimensions[get_column_letter(idx)].width = 18
        validations = {
            "묶음": '"A,B,C,D,E"',
            "유형": '"T1 정확표기,T2 동의어,T3 개념확장,T4 구어·문장,T5 서비스"',
            "작성방식": '"직접,LLM후보"',
        }
        for col_name, formula in validations.items():
            letter = get_column_letter(INPUT_COLUMNS.index(col_name) + 1)
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
            dv.add(f"{letter}2:{letter}{INPUT_ROWS + 1}")
            ws.add_data_validation(dv)
    logger.info("저장: %s", output)


if __name__ == "__main__":
    main()
