"""서비스셋 판정 pool을 만든다 — 쿼리 가정위치 반경 안 가맹점만 후보로(docs/SERVICE_ANNOTATION_GUIDELINE.md 4절).

단계(설정 `judging` 블록, configs/query/service_queryset_v1.yaml)
  1) export  임베딩 run은 Colab에서 만든다(2026-10-20 결정). 반경 안 가맹점·쿼리를 colab_upload/service_eval/로 내보낸다
             → Drive store-search-ai/service_eval/에 올리고 colab/service_dense_runs.ipynb 실행 → 나온 <모델>.csv를
             <benchmark_dir>/runs/에 넣는다.
     (encode  로컬 CPU로 임베딩 캐시를 만드는 대안. 모델당 1시간 안팎)
  2) runs    char BM25 · 키워드 후보 run(로컬)을 반경 안 상위 run_depth개까지 만든다 → <benchmark_dir>/runs/.
             로컬 임베딩 캐시가 있으면 그 run도 만들고, 없으면 Colab에서 받은 run 파일을 그대로 쓴다.
  3) pool    depth별 판정 건수 표(--dry-run) → depth를 정해 판정 시트 생성 → <benchmark_dir>/annotations/<round>/

사용법:
    python scripts/build_service_pool.py export                     # Colab 입력
    python scripts/build_service_pool.py encode                     # (대안) 로컬 임베딩 캐시
    python scripts/build_service_pool.py encode --model configs/models/bge_m3.yaml
    python scripts/build_service_pool.py runs
    python scripts/build_service_pool.py pool --dry-run
    python scripts/build_service_pool.py pool --round service_v1 --depth 10
    # anchor 사례(정답 없는 쿼리에 앵커 근처 위치 추가): 같은 단계를 --set anchor로
    python scripts/build_service_pool.py export --set anchor   # → colab_upload/service_eval_anchor/
    python scripts/build_service_pool.py runs --set anchor     # Colab run은 <benchmark_dir>/runs_anchor/에
    python scripts/build_service_pool.py pool --set anchor --round service_v2 --depth 10
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
from datetime import UTC, datetime
from pathlib import Path

import numpy as np
import pandas as pd

from store_search_ai.data.service_pool import (
    RUN_COLUMNS,
    annotation_sheet,
    assign_annotators,
    candidates_within,
    distances_km,
    double_annotation_queries,
    keyword_run,
    lexical_query,
    pool_pairs,
    runs_outside_radius,
    top_k_run,
)
from store_search_ai.pipeline.common import load_config, sha256_file
from store_search_ai.retrieval.lexical import build_pool_document_texts, fit_bm25_scorer

logger = logging.getLogger(__name__)

CACHE_DIR = Path("data/interim/service_eval")
DOC_TEXT_COLUMN = "search_text_t1_minimal"   # 모델 평가와 같은 문서 텍스트(가맹점명 + 취급품목)


SETS = {
    # 평가 사례 묶음 → (Colab 입력 폴더, run 폴더 이름). 이미 판정한 묶음의 run은 다시 만들지 않는다 — BM25 색인 범위가 바뀌면
    # 순위가 흔들려 판정 범위와 어긋나기 때문이다. 새 쿼리는 새 묶음(extra)으로 따로 만든다.
    "main": (Path("colab_upload/service_eval"), "runs"),
    "anchor": (Path("colab_upload/service_eval_anchor"), "runs_anchor"),
    "extra": (Path("colab_upload/service_eval_extra"), "runs_extra"),
    "extra_b": (Path("colab_upload/service_eval_extra_b"), "runs_extra_b"),   # 추가 B(2026-10-26, 주 평가셋 비율)
    "anchor_extra": (Path("colab_upload/service_eval_anchor_extra"), "runs_anchor_extra"),   # 추가 쿼리의 anchor(2026-10-26)
}


def load_inputs(cfg: dict, bench: dict, set_name: str = "main") -> tuple[pd.DataFrame, pd.DataFrame, dict[str, np.ndarray]]:
    """평가 사례 표(query_id = 사례 ID, base_query_id = 판정에 쓰는 쿼리 ID), 코퍼스, 사례별 반경 안 후보.

    main: 활성 쿼리(사례 ID = 쿼리 ID). 나머지: instances.csv에서 batch가 묶음 이름과 같은 사례
    (anchor = 처음 anchor 사례, anchor_extra = 나중에 더한 쿼리의 anchor, extra·extra_b = 나중에 더한 쿼리).
    """

    if set_name != "main":
        inst = pd.read_csv(Path(cfg["judging"]["benchmark_dir"]) / "instances.csv", encoding="utf-8-sig", keep_default_na=False)
        if "batch" not in inst:
            inst["batch"] = inst["context"].map({"random": "v1", "anchor": "anchor"})
        pick = inst["batch"] == set_name
        queries = inst[pick].rename(columns={"query_id": "base_query_id"})
        queries = queries.rename(columns={"instance_id": "query_id"}).reset_index(drop=True)
    else:   # main: 처음 판정한 v1 쿼리(benchmark_dir/queries.csv가 있으면 그 목록, 없으면 가져온 활성 쿼리 전부)
        queries = pd.read_csv(cfg["import"]["output"], encoding="utf-8-sig", keep_default_na=False)
        queries = queries[queries["status"] == "active"].reset_index(drop=True)
        v1_list = Path(cfg["judging"]["benchmark_dir"]) / "queries.csv"
        if v1_list.exists():
            v1 = set(pd.read_csv(v1_list, encoding="utf-8-sig", keep_default_na=False)["query_id"])
            queries =queries[queries["query_id"].isin(v1)].reset_index(drop=True)
        queries = queries.assign(base_query_id=queries["query_id"])
    corpus = pd.read_parquet(bench["corpus_path"], columns=[
        "doc_id", "store_name", "item_text", "market_name", "latitude", "longitude", "geo_status", DOC_TEXT_COLUMN])
    corpus = corpus[corpus["geo_status"] == "VALID"].reset_index(drop=True)
    cands = candidates_within(queries, corpus, cfg["judging"]["judge_radius_km"])
    return queries, corpus, cands


def union_index(cands: dict[str, np.ndarray]) -> np.ndarray:
    return np.unique(np.concatenate(list(cands.values())))


def cache_paths(name: str) -> tuple[Path, Path]:
    return CACHE_DIR / f"emb_{name}.npy", CACHE_DIR / f"emb_{name}.doc_ids.npy"


def encode(args, cfg: dict, bench: dict) -> None:
    from store_search_ai.models.sentence_transformer_encoder import (
        SentenceTransformerEncoder,
    )

    _, corpus, cands = load_inputs(cfg, bench)   # 로컬 캐시는 main 사례만
    idx = union_index(cands)
    doc_ids = corpus["doc_id"].to_numpy()[idx]
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    for model_cfg in ([args.model] if args.model else cfg["judging"]["dense_models"]):
        mcfg = load_config(model_cfg)
        emb_path, ids_path = cache_paths(mcfg["name"])
        if emb_path.exists() and ids_path.exists() and np.array_equal(np.load(ids_path, allow_pickle=True), doc_ids):
            logger.info("%s: 캐시가 이미 있음 — 건너뜀", mcfg["name"])
            continue
        logger.info("%s: 반경 안 가맹점 %d개 인코딩", mcfg["name"], len(idx))
        encoder = SentenceTransformerEncoder(mcfg)
        emb = encoder.encode_corpus(corpus[DOC_TEXT_COLUMN].to_numpy()[idx].tolist())
        np.save(emb_path, emb)
        np.save(ids_path, doc_ids.astype(object), allow_pickle=True)
        logger.info("저장: %s", emb_path)


def export(args, cfg: dict, bench: dict) -> None:
    """Colab(colab/service_dense_runs.ipynb) 입력: 판정 반경 안에 드는 가맹점과 활성 쿼리, 반경·depth."""

    jcfg = cfg["judging"]
    queries, corpus, cands = load_inputs(cfg, bench, args.set)
    idx = union_index(cands)
    out = SETS[args.set][0]
    out.mkdir(parents=True, exist_ok=True)
    docs = corpus.iloc[idx][["doc_id", DOC_TEXT_COLUMN, "latitude", "longitude"]].rename(columns={DOC_TEXT_COLUMN: "text"})
    docs.to_csv(out / "docs.csv", index=False, encoding="utf-8-sig", lineterminator="\n")
    queries[["query_id", "query", "latitude", "longitude"]].to_csv(out / "queries.csv", index=False, encoding="utf-8-sig",
                                                                  lineterminator="\n")
    meta = {"judge_radius_km": jcfg["judge_radius_km"], "run_depth": jcfg["run_depth"],
            "report_radii_km": sorted(set(jcfg["report_radii_km"])), "run_dir": SETS[args.set][1], "docs": len(docs),
            "queries": len(queries), "doc_text": DOC_TEXT_COLUMN, "corpus_path": bench["corpus_path"]}
    (out / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    logger.info("가맹점 %d개, 사례 %d개 → %s (Drive store-search-ai/%s/에 올림)", len(docs), len(queries), out, out.name)


def runs(args, cfg: dict, bench: dict) -> None:
    from store_search_ai.models.sentence_transformer_encoder import (
        SentenceTransformerEncoder,
    )

    jcfg = cfg["judging"]
    queries, corpus, cands = load_inputs(cfg, bench, args.set)
    idx = union_index(cands)
    doc_ids = corpus["doc_id"].to_numpy()
    run_dir = Path(jcfg["benchmark_dir"]) / SETS[args.set][1]
    run_dir.mkdir(parents=True, exist_ok=True)
    depth = jcfg["run_depth"]
    radii = sorted(set(jcfg["report_radii_km"]))   # 반경마다 상위 depth개(작은 반경 순위가 잘리지 않게)
    dist = distances_km(queries, corpus, cands)

    # BM25: 판정 반경 안에 드는 가맹점만으로 색인(IDF도 그 범위 기준)
    texts = build_pool_document_texts(corpus).to_numpy()
    pos = np.full(len(corpus), -1)
    pos[idx] = np.arange(len(idx))
    logger.info("char BM25 색인: %d개", len(idx))
    bm25 = fit_bm25_scorer(pd.Series(texts[idx]))
    frames = [top_k_run(qid, cands[qid], bm25(lexical_query(q, region))[pos[cands[qid]]], doc_ids, depth, "char_bm25",
                        dist[qid], radii)
              for qid, q, region in zip(queries["query_id"], queries["query"], queries["region_text"], strict=True)]
    write_run(pd.concat(frames, ignore_index=True), run_dir / "char_bm25.csv")

    frames = [keyword_run(qid, lexical_query(q, region), cands[qid], texts, doc_ids, jcfg["keyword_per_query"], dist[qid],
                          radii)
              for qid, q, region in zip(queries["query_id"], queries["query"], queries["region_text"], strict=True)]
    write_run(pd.concat(frames, ignore_index=True), run_dir / "keyword.csv")

    for model_cfg in jcfg["dense_models"]:
        mcfg = load_config(model_cfg)
        emb_path, ids_path = cache_paths(mcfg["name"])
        if args.set != "main" or not emb_path.exists():
            colab_run = run_dir / f"{mcfg['name']}.csv"
            if colab_run.exists():
                logger.info("%s: 로컬 캐시 없음 — Colab에서 받은 run(%s)을 그대로 씀", mcfg["name"], colab_run)
            else:
                logger.warning("%s: run이 없음 — Colab(colab/service_dense_runs.ipynb)에서 만들어 %s에 넣을 것", mcfg["name"], colab_run)
            continue
        cached_ids = np.load(ids_path, allow_pickle=True)
        if not np.array_equal(cached_ids, doc_ids[idx]):
            raise SystemExit(f"{mcfg['name']}: 캐시의 가맹점 목록이 지금 쿼리·반경과 다릅니다 — encode를 다시 실행하세요.")
        emb = np.load(emb_path)
        q_emb = SentenceTransformerEncoder(mcfg).encode_queries(queries["query"].tolist())
        frames = [top_k_run(qid, cands[qid], emb[pos[cands[qid]]] @ q_emb[i], doc_ids, depth, mcfg["name"], dist[qid], radii)
                  for i, qid in enumerate(queries["query_id"])]
        write_run(pd.concat(frames, ignore_index=True), run_dir / f"{mcfg['name']}.csv")


def write_run(run: pd.DataFrame, path: Path) -> None:
    run[RUN_COLUMNS].to_csv(path, index=False, encoding="utf-8-sig", lineterminator="\n")
    logger.info("저장: %s (%d행, 쿼리 %d개)", path, len(run), run["query_id"].nunique())


def pool(args, cfg: dict, bench: dict) -> None:
    jcfg = cfg["judging"]
    bdir = Path(jcfg["benchmark_dir"])
    run_files = sorted((bdir / SETS[args.set][1]).glob("*.csv"))
    if not run_files:
        raise SystemExit("run이 없습니다 — 먼저 runs")
    all_runs = pd.concat([pd.read_csv(p, encoding="utf-8-sig") for p in run_files], ignore_index=True)
    queries, corpus, _cands = load_inputs(cfg, bench, args.set)
    stale = {k: v for k, v in runs_outside_radius(all_runs, queries, corpus, jcfg["judge_radius_km"]).items() if v}
    if stale:
        raise SystemExit(f"반경 밖이거나 지금 쿼리에 없는 run 행이 있습니다 {stale} — 쿼리 위치가 바뀐 뒤의 옛 run입니다. "
                         "export → Colab → runs를 다시 하세요.")

    judged = previously_judged(bdir)
    radii = sorted(set(jcfg["report_radii_km"]))   # 채점하는 반경마다 상위 depth개를 판정
    rows = []
    for d in sorted({5, 10, 20, args.depth or jcfg["depth"]}):
        pairs = judging_pairs(all_runs, queries, d, judged, corpus, radii)
        per_q = pairs.groupby("query_id").size()
        rows.append({"depth": d, "새 판정 쌍": len(pairs), "쿼리 수": len(per_q),
                     "쿼리당 평균": round(per_q.mean(), 1) if len(per_q) else 0, "쿼리당 최대": int(per_q.max()) if len(per_q) else 0})
    print(f"========== 판정 건수(시스템 {sorted(all_runs['system'].unique())}, 쿼리 {len(queries)}개) ==========")
    print(pd.DataFrame(rows).to_string(index=False))
    share = jcfg["double_annotation_share"]
    print(f"(쿼리의 {share:.0%}는 A·B 두 명이 판정 — 사람 판정 건수는 그만큼 더 많음)")
    if args.dry_run:
        return
    if not args.round:
        raise SystemExit("--round가 필요합니다.")

    round_dir = bdir / "annotations" / args.round
    if (round_dir / "pool_manifest.json").exists():
        raise SystemExit(f"{round_dir}에 이미 라운드가 있습니다 — 채우는 중인 시트를 덮어쓰지 않도록 다른 --round를 쓰세요.")
    depth = args.depth or jcfg["depth"]
    pairs = judging_pairs(all_runs, queries, depth, judged, corpus, radii)
    seed = cfg["random_seed"]
    base = (pd.read_csv(cfg["import"]["output"], encoding="utf-8-sig", keep_default_na=False)
            if args.set != "main" else queries)   # 판정 시트의 쿼리 정보(문자열·묶음·의도)
    sheet = annotation_sheet(pairs, base, corpus, seed)
    double = double_annotation_queries(sorted(pairs["query_id"].unique()), share, seed)
    writers = {qid: set(str(w).split("|")) for qid, w in zip(base["query_id"], base["writers"], strict=True)}
    assignment = assign_annotators(pairs.groupby("query_id").size(), double, jcfg["annotators"], writers, seed)
    round_dir.mkdir(parents=True, exist_ok=True)
    for name, qids in assignment.items():
        part = sheet[sheet["query_id"].isin(qids)].reset_index(drop=True)
        part.to_csv(round_dir / f"annotation_{name}.csv", index=False, encoding="utf-8-sig", lineterminator="\n")
        write_annotation_xlsx(part, round_dir / f"판정_{name}.xlsx")
        logger.info("판정자 %s: 쿼리 %d개(이중 판정 %d 포함), %d쌍", name, len(qids), len(double), len(part))
    pairs.to_csv(round_dir / "pool.csv", index=False, encoding="utf-8-sig", lineterminator="\n")
    if args.set == "main":
        queries.drop(columns=["base_query_id"]).to_csv(bdir / "queries.csv", index=False, encoding="utf-8-sig",
                                                       lineterminator="\n")
    manifest = {
        "round": args.round,
        "set": args.set,
        "reused_judgments": "판정은 (쿼리, 가맹점) 단위·위치 무관 — 앞 라운드에서 판정한 쌍은 다시 판정하지 않음",
        "created_at": datetime.now(UTC).isoformat(),
        "judge_radius_km": jcfg["judge_radius_km"],
        "report_radii_km": jcfg["report_radii_km"],
        "depth": depth,
        "depth_radii_km": radii,
        "queries": int(pairs["query_id"].nunique()),
        "pairs": len(pairs),
        "double_annotation_queries": len(double),
        "double_annotation_pairs": int(sheet["query_id"].isin(double).sum()),
        "annotators": {name: {"writer_id": jcfg["annotators"][name], "queries": len(qids),
                              "pairs": int(sheet["query_id"].isin(qids).sum())} for name, qids in assignment.items()},
        "runs_sha256": {p.name: sha256_file(p) for p in run_files},
        "queries_sha256": hashlib.sha256((bdir / "queries.csv").read_bytes()).hexdigest(),
        "instances_sha256": sha256_file(bdir / "instances.csv") if (bdir / "instances.csv").exists() else None,
        "guideline": "docs/SERVICE_ANNOTATION_GUIDELINE.md + benchmark/storesearch_ko_v1/annotation_guideline.md(v1.2)",
    }
    (round_dir / "pool_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    logger.info("저장: %s (전체 %d쌍, 이중 판정 %d쌍)", round_dir, len(sheet), manifest["double_annotation_pairs"])


def previously_judged(bdir: Path) -> set[tuple[str, str]]:
    """앞 라운드에서 판정 대상이었던 (쿼리, 가맹점). 판정은 위치와 무관하므로 다른 위치 사례에서 다시 나와도 다시 판정하지 않는다."""

    out: set[tuple[str, str]] = set()
    for path in (bdir / "annotations").glob("*/pool.csv"):
        p = pd.read_csv(path, encoding="utf-8-sig")
        out |= set(zip(p["query_id"], p["doc_id"], strict=True))
    return out


def within_radius(runs: pd.DataFrame, queries: pd.DataFrame, corpus: pd.DataFrame, radius_km: float) -> pd.DataFrame:
    """run에서 사례 위치로부터 radius_km 안만 남기고 순위를 다시 매긴다(1km 채점에 쓰는 순위와 같음)."""

    loc = queries.set_index("query_id")[["latitude", "longitude"]]
    c = corpus.set_index("doc_id")[["latitude", "longitude"]]
    r = runs[runs["query_id"].isin(loc.index)].copy()
    a = np.radians(loc.loc[r["query_id"]].to_numpy(float))
    b = np.radians(c.loc[r["doc_id"]].to_numpy(float))
    h = np.sin((b[:, 0] - a[:, 0]) / 2) ** 2 + np.cos(a[:, 0]) * np.cos(b[:, 0]) * np.sin((b[:, 1] - a[:, 1]) / 2) ** 2
    r = r[2 * 6371.0 * np.arcsin(np.sqrt(h)) <= radius_km + 1e-9].sort_values(["system", "query_id", "rank"])
    return r.assign(rank=r.groupby(["system", "query_id"]).cumcount() + 1)


def judging_pairs(runs: pd.DataFrame, queries: pd.DataFrame, depth: int, judged: set[tuple[str, str]],
                  corpus: pd.DataFrame | None = None, radii: list[float] | None = None) -> pd.DataFrame:
    """사례별 run 상위 depth개(반경마다 다시 매긴 순위로, 합집합) + 앵커 가게 → 판정 쿼리 ID(base) 기준 (쿼리, 가맹점).
    앞 라운드에서 판정한 쌍은 뺀다. 채점 반경이 둘(1km·2km)이면 1km 안 상위 depth개도 판정해야 1km 지표의 판정 누락이 없다."""

    if radii and corpus is not None:
        runs = pd.concat([within_radius(runs, queries, corpus, r) for r in radii], ignore_index=True)
    pairs = pool_pairs(runs, depth).rename(columns={"query_id": "instance_id"})
    to_base = dict(zip(queries["query_id"], queries["base_query_id"], strict=True))
    pairs["query_id"] = pairs["instance_id"].map(to_base)
    if "anchor_doc_id" in queries:
        anchors = queries.loc[queries["anchor_doc_id"] != "", ["query_id", "base_query_id", "anchor_doc_id"]]
        pairs = pd.concat([pairs, pd.DataFrame({"instance_id": anchors["query_id"], "doc_id": anchors["anchor_doc_id"],
                                                "pool_sources": "anchor", "query_id": anchors["base_query_id"]})])
    pairs = (pairs.groupby(["query_id", "doc_id"], as_index=False)
             .agg(pool_sources=("pool_sources", lambda s: "|".join(sorted(set("|".join(s).split("|"))))),
                  instance_ids=("instance_id", lambda s: "|".join(sorted(set(s))))))
    keep = [(q, d) not in judged for q, d in zip(pairs["query_id"], pairs["doc_id"], strict=True)]
    return pairs[keep].reset_index(drop=True)


def write_annotation_xlsx(sheet: pd.DataFrame, path: Path) -> None:
    """판정자용 엑셀: relevance(0~3)·uncertain(Y) 드롭다운, 채울 칸 색칠, 첫 행 고정. 회수는 이 파일이나 같은 열의 CSV로 받는다."""

    from openpyxl.styles import Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    widths = {"query": 22, "intent_hint": 22, "store_name": 26, "item_text": 30, "market_name": 22, "annotator_note": 30}
    fill = PatternFill("solid", fgColor="FFF7D6")
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        sheet.to_excel(writer, sheet_name="판정", index=False)
        ws = writer.sheets["판정"]
        ws.freeze_panes = "C2"
        cols = list(sheet.columns)
        for i, col in enumerate(cols, start=1):
            letter = get_column_letter(i)
            ws.column_dimensions[letter].width = widths.get(col, 10)
            if col in ("relevance", "uncertain", "annotator_note"):
                for r in range(2, len(sheet) + 2):
                    ws[f"{letter}{r}"].fill = fill
            if col in ("item_text", "intent_hint"):
                for r in range(2, len(sheet) + 2):
                    ws[f"{letter}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        for col, values in (("relevance", '"0,1,2,3"'), ("uncertain", '"Y"')):
            letter = get_column_letter(cols.index(col) + 1)
            dv = DataValidation(type="list", formula1=values, allow_blank=True)
            dv.add(f"{letter}2:{letter}{len(sheet) + 1}")
            ws.add_data_validation(dv)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["export", "encode", "runs", "pool"])
    parser.add_argument("--config", default="configs/query/service_queryset_v1.yaml")
    parser.add_argument("--benchmark-config", default="configs/benchmark/storesearch_ko_v1.yaml")
    parser.add_argument("--model", help="encode할 모델 yaml 하나(기본: 설정의 dense_models 전부)")
    parser.add_argument("--round", help="pool: 판정 라운드 이름(예: service_v1)")
    parser.add_argument("--depth", type=int, help="pool: 시스템마다 상위 몇 개(기본: 설정 judging.depth)")
    parser.add_argument("--dry-run", action="store_true", help="pool: 판정 건수 표만")
    parser.add_argument("--set", choices=sorted(SETS), default="main",
                        help="평가 사례 묶음: main(활성 쿼리의 기본 위치) / anchor(정답 없는 쿼리의 앵커 근처 위치, instances.csv)")
    args = parser.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(message)s")
    cfg = load_config(args.config)
    bench = load_config(args.benchmark_config)
    {"export": export, "encode": encode, "runs": runs, "pool": pool}[args.command](args, cfg, bench)


if __name__ == "__main__":
    main()
